"""Generate the documented 40-row donor upload workbook."""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from app.seed.donors import sample_invalid_rows, sample_valid_rows

SAMPLE_WORKBOOK_PATH = Path(__file__).resolve().parent / "assets" / "sample-donors.xlsx"
SAMPLE_HEADERS = (
    "name",
    "phone",
    "segment",
    "language",
    "city",
    "blood_group",
    "last_donation_date",
)


def generate_sample_workbook(path: Path = SAMPLE_WORKBOOK_PATH) -> Path:
    """Write the deterministic sample upload with 35 valid and five invalid rows."""
    workbook = Workbook()
    sheet = workbook.active
    if sheet is None:
        raise RuntimeError("Workbook did not create an active sheet")
    sheet.title = "Donors"
    sheet.append(SAMPLE_HEADERS)
    for row in (*sample_valid_rows(), *sample_invalid_rows()):
        sheet.append(tuple(row[header] for header in SAMPLE_HEADERS))
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(fill_type="solid", fgColor="005B96")
    sheet.freeze_panes = "A2"
    widths = (24, 18, 15, 12, 20, 14, 20)
    for column_index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(column_index)].width = width
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    return path
