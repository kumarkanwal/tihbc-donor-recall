"""Tests for the generated donor-upload workbook asset."""

from pathlib import Path

from openpyxl import load_workbook

from app.seed.sample_workbook import SAMPLE_HEADERS, generate_sample_workbook


def test_sample_workbook_has_40_rows_and_five_documented_invalid_examples(
    tmp_path: Path,
) -> None:
    path = generate_sample_workbook(tmp_path / "sample-donors.xlsx")
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.active
    assert sheet is not None
    rows = list(sheet.iter_rows(values_only=True))

    assert rows[0] == SAMPLE_HEADERS
    assert len(rows[1:]) == 40
    assert rows[36][1] == "12345"
    assert rows[37][2] == "vip"
    assert rows[38][0] is None
    assert rows[39][1] == rows[1][1]
    assert rows[40][6] == "2099-01-01"
